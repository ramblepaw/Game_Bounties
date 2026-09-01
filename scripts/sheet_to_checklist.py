"""Convert a Google Sheets-exported .xlsx tracker into the app's checklist import JSON.

Usage:
    py scripts/sheet_to_checklist.py "Checklists/Wind Waker Checklist.xlsx"

Every tracker is shaped around its own game, so this script does not try to
guess what a block of checkboxes means. Run it once to see the blocks it found,
then re-run with a --block rule for any that should become something other than
plain checkboxes.

Structure comes from the fill colors, which is how these trackers mark
sections: a lone saturated cell is a section header, and the lighter fill
beneath it is that section's body. Both carry over to the module's
titleBgColor/bgColor.

A grid of checkboxes in a sheet is almost always a workaround for a cell only
being able to hold one checkmark -- the app has richer item kinds, so the
mapping usually dissolves the grid rather than preserving it:

    module-per-row  each row label becomes its own module
    stage           each row becomes one STAGE item; columns are the stages
    counter         each row becomes one COUNTER item; target = its checkboxes
    items           (default) one checkbox per box

The import creates a brand-new checklist, so a bad run is undone by deleting it.
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from dataclasses import dataclass
from pathlib import Path

try:
    import openpyxl
except ImportError:  # pragma: no cover
    sys.exit("openpyxl is required:  py -m pip install openpyxl")


MAX_GRID_COLUMNS = 12

# Sheets that duplicate something the app already tracks natively.
DEFAULT_SKIP = ["Overview", "Tracker", "Playtime Log"]

# A fill used on at most this many cells reads as deliberate header styling
# rather than a body wash. Every header in the sample workbook sits at 1-4.
HEADER_FILL_MAX_CELLS = 4


# --------------------------------------------------------------------------
# workbook reading
# --------------------------------------------------------------------------

def argb_to_hex(argb):
    """FFD9EAD3 -> #D9EAD3. Returns None for unset/auto fills."""
    if not argb or not isinstance(argb, str):
        return None
    if len(argb) == 8:
        argb = argb[2:]
    if len(argb) != 6 or argb == "000000":
        return None
    return "#" + argb.upper()


def cell_fill(cell):
    fill = cell.fill
    if not fill or not fill.patternType:
        return None
    rgb = getattr(fill.start_color, "rgb", None)
    if rgb in (None, "00000000"):
        return None
    return rgb if isinstance(rgb, str) else None


def is_bool(value):
    return isinstance(value, bool)


def is_text(value):
    # Formula strings (an Overview-style rollup that leaked into a data sheet)
    # are not item titles.
    return isinstance(value, str) and bool(value.strip()) and not value.startswith("=")


def clean(value):
    return re.sub(r"\s+", " ", str(value)).strip()


@dataclass
class Block:
    """One header'd region of a sheet, before it is mapped to anything."""
    name: str
    title_fill: str | None
    row_start: int
    row_end: int
    col_start: int
    col_end: int
    body_fill: str | None = None


def find_header_cells(ws, fill_counts):
    headers = []
    for row in ws.iter_rows():
        for cell in row:
            if not is_text(cell.value):
                continue
            fill = cell_fill(cell)
            if fill and fill_counts.get(fill, 0) <= HEADER_FILL_MAX_CELLS:
                headers.append((cell.row, cell.column, clean(cell.value), fill))
    return sorted(headers)


def find_blocks(ws, headers):
    """Carve the sheet into (rows x cols) windows, one per header.

    Headers sharing a row sit side by side (two trackers in adjacent column
    blocks), so a row's window is split at each following header's column.
    """
    blocks = []
    header_rows = sorted({r for r, _, _, _ in headers})

    for row, col, text, fill in headers:
        next_header_row = next((r for r in header_rows if r > row), None)
        row_end = (next_header_row - 1) if next_header_row else ws.max_row
        same_row_cols = sorted(c for r, c, _, _ in headers if r == row and c > col)
        col_end = (same_row_cols[0] - 1) if same_row_cols else ws.max_column
        blocks.append(Block(text, fill, row + 1, row_end, col, col_end))

    # Anything above the first header is still real data -- most sheets open
    # with a title in A1 whose fill matches the body wash, so it is not
    # detected as a header but still names the block below it.
    first_header_row = header_rows[0] if headers else ws.max_row + 1
    if first_header_row > 2:
        title = ws.cell(row=1, column=1).value
        named = is_text(title)
        blocks.insert(0, Block(clean(title) if named else clean(ws.title),
                               cell_fill(ws.cell(row=1, column=1)),
                               2 if named else 1, first_header_row - 1,
                               1, ws.max_column))
    elif not headers:
        blocks.append(Block(clean(ws.title), None, 1, ws.max_row, 1, ws.max_column))

    return blocks


def block_cells(ws, blk):
    for r in range(blk.row_start, blk.row_end + 1):
        for c in range(blk.col_start, blk.col_end + 1):
            yield ws.cell(row=r, column=c)


def dominant_body_fill(ws, blk):
    counts = {}
    for cell in block_cells(ws, blk):
        fill = cell_fill(cell)
        if fill and fill != blk.title_fill:
            counts[fill] = counts.get(fill, 0) + 1
    return max(counts, key=counts.get) if counts else None


def read_block(ws, blk):
    """Read a block's checkboxes into either a flat list or a labelled grid.

    Returns ("grid", [(row_label, [col_label, ...]), ...]) when the block is a
    matrix, or ("list", [title, ...]) otherwise. Ragged rows are preserved:
    a row only reports the columns that actually hold a checkbox.
    """
    boxes = [c for c in block_cells(ws, blk) if is_bool(c.value)]
    if not boxes:
        return None, []

    box_rows = {c.row for c in boxes}
    rows = range(blk.row_start, blk.row_end + 1)
    cols = range(blk.col_start, blk.col_end + 1)

    def row_texts(r):
        return [c for c in (ws.cell(row=r, column=cc) for cc in cols) if is_text(c.value)]

    head = blk.row_start
    head_labels = row_texts(head)

    # --- grid: labels across the top, labels down the left, checkboxes inside
    if (
        head not in box_rows
        and len(head_labels) >= 2
        and any(is_text(ws.cell(row=r, column=blk.col_start).value) for r in rows if r > head)
    ):
        col_labels = {c.column: clean(c.value) for c in head_labels}
        grid = []
        for r in rows:
            if r == head:
                continue
            row_label = ws.cell(row=r, column=blk.col_start).value
            if not is_text(row_label):
                continue
            labels = [col_labels.get(c, "") for c in cols if is_bool(ws.cell(row=r, column=c).value)]
            if labels:
                grid.append((clean(row_label), labels))
        if grid:
            return "grid", grid

    # --- stacked: a row of labels, the row of checkboxes beneath it, repeating
    label_rows = [r for r in rows if r not in box_rows and row_texts(r)]
    if label_rows and all((r + 1) in box_rows for r in label_rows):
        titles = [
            clean(ws.cell(row=r, column=c).value)
            for r in label_rows
            for c in cols
            if is_bool(ws.cell(row=r + 1, column=c).value)
            and is_text(ws.cell(row=r, column=c).value)
        ]
        if titles:
            return "list", titles

    # --- pairs: a label with its checkbox in the next column over
    titles = []
    for cell in sorted(boxes, key=lambda c: (c.row, c.column)):
        left = ws.cell(row=cell.row, column=cell.column - 1) if cell.column > 1 else None
        if left is not None and is_text(left.value):
            titles.append(clean(left.value))
        else:
            header = ws.cell(row=blk.row_start, column=cell.column)
            titles.append(clean(header.value) if is_text(header.value) else "Row " + str(cell.row))
    return "list", titles


# --------------------------------------------------------------------------
# mapping a block onto the app's item kinds
# --------------------------------------------------------------------------

def make_module(name, blk, items, grid_columns=4):
    return {
        "name": name,
        "itemLayout": "GRID",
        "gridColumns": max(2, min(grid_columns, MAX_GRID_COLUMNS)),
        "span": 4,
        "titleBgColor": argb_to_hex(blk.title_fill),
        "bgColor": argb_to_hex(blk.body_fill),
        "stages": [],
        "items": items,
    }


def map_block(blk, shape, data, rule):
    """Turn one block into one or more modules, per its mapping rule."""
    mode = rule.get("mode", "items")

    if shape == "grid" and mode == "module-per-row":
        # The row label is the module; repeating it in every item name just
        # re-adds the constraint the grid existed to work around.
        return [
            make_module(row_label, blk, [{"title": lbl or row_label} for lbl in labels],
                        grid_columns=len(labels))
            for row_label, labels in data
        ]

    if shape == "grid" and mode == "stage":
        # Columns are not separate objectives -- they are progressive states
        # of one. targetCount is left unset: the app resolves a STAGE item's
        # weight from its module's stage list.
        stage_names = rule.get("stages") or list(dict.fromkeys(
            lbl for _, labels in data for lbl in labels if lbl
        ))
        module = make_module(blk.name, blk, [{"title": r, "kind": "STAGE"} for r, _ in data])
        module["stages"] = [{"name": n} for n in stage_names]
        return [module]

    if shape == "grid" and mode == "counter":
        return [make_module(blk.name, blk, [
            {"title": row_label, "kind": "COUNTER", "targetCount": len(labels)}
            if len(labels) > 1 else {"title": row_label}
            for row_label, labels in data
        ])]

    if shape == "grid":
        return [make_module(blk.name, blk, [
            {"title": (row_label + " " + lbl).strip()}
            for row_label, labels in data for lbl in labels
        ])]

    return [make_module(blk.name, blk, [{"title": t} for t in data])]


def convert_sheet(ws, config):
    fill_counts = {}
    for row in ws.iter_rows():
        for cell in row:
            fill = cell_fill(cell)
            if fill:
                fill_counts[fill] = fill_counts.get(fill, 0) + 1

    rules = config.get("blocks", {})
    url_template = config.get("urlTemplate")
    modules = []

    for blk in find_blocks(ws, find_header_cells(ws, fill_counts)):
        shape, data = read_block(ws, blk)
        if not shape:
            continue
        blk.body_fill = dominant_body_fill(ws, blk)
        rule = rules.get(ws.title.strip() + "::" + blk.name) or rules.get(blk.name) or {}
        modules.extend(map_block(blk, shape, data, rule))

    for order, module in enumerate(modules):
        module["order"] = order
        for i, item in enumerate(module["items"]):
            item["order"] = i
            item.setdefault("kind", "CHECKBOX")
            if url_template and item["kind"] not in ("TITLE",):
                slug = re.sub(r"[^A-Za-z0-9]+", "_", item["title"]).strip("_")
                item["url"] = url_template.format(title=item["title"], slug=slug)

    if not modules:
        return None
    return {"title": clean(ws.title), "sections": modules}


# --------------------------------------------------------------------------

def main():
    ap = argparse.ArgumentParser(description=__doc__,
                                 formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("workbook", type=Path)
    ap.add_argument("-o", "--out", type=Path, help="output JSON (default: alongside the workbook)")
    ap.add_argument("-n", "--name", help="checklist name (default: the workbook filename)")
    ap.add_argument("--skip", nargs="*", default=DEFAULT_SKIP,
                    help="sheets to ignore (default: " + " ".join(DEFAULT_SKIP) + ")")
    ap.add_argument("--url-template", help="per-item link, e.g. 'https://zeldawiki.wiki/wiki/{slug}'")
    ap.add_argument("--block", action="append", default=[], metavar="NAME=MODE",
                    help="how one block converts, e.g. 'Dungeon Items=module-per-row' or "
                         "'Treasure Charts=stage:Acquired,Found'. Repeatable; anything "
                         "unmentioned becomes plain checkboxes.")
    args = ap.parse_args()

    if not args.workbook.exists():
        return "No such workbook: " + str(args.workbook)

    blocks = {}
    for spec in args.block:
        if "=" not in spec:
            return "Bad --block (expected NAME=MODE): " + spec
        name, _, mode = spec.partition("=")
        mode, _, stages = mode.partition(":")
        if mode not in ("items", "module-per-row", "stage", "counter"):
            return "Unknown mode {!r} in --block {!r}".format(mode, spec)
        rule = {"mode": mode}
        if stages:
            rule["stages"] = [s.strip() for s in stages.split(",") if s.strip()]
        blocks[name.strip()] = rule

    config = {"blocks": blocks, "urlTemplate": args.url_template}

    wb = openpyxl.load_workbook(args.workbook, data_only=True)
    skip = {s.strip().casefold() for s in args.skip}

    tabs, palette = [], {}
    for ws in wb.worksheets:
        if ws.title.strip().casefold() in skip:
            print("  skip   " + repr(ws.title))
            continue
        tab = convert_sheet(ws, config)
        if not tab:
            print("  empty  " + repr(ws.title) + " -- no checkboxes found")
            continue
        tab["order"] = len(tabs)
        tabs.append(tab)
        n = sum(len(s["items"]) for s in tab["sections"])
        print("  tab    {!r}: {} modules, {} items".format(tab["title"], len(tab["sections"]), n))
        for s in tab["sections"]:
            kinds = {i["kind"] for i in s["items"]} - {"CHECKBOX"}
            note = (" [" + ",".join(sorted(kinds)) + "]") if kinds else ""
            print("           - {!r} ({}){}".format(s["name"], len(s["items"]), note))
            for key in ("titleBgColor", "bgColor"):
                if s[key]:
                    palette.setdefault(s[key], (tab["title"] + " " + s["name"])[:40])

    if not tabs:
        return "Nothing to import -- every sheet was skipped or had no checkboxes."

    checklist = {
        "name": args.name or args.workbook.stem,
        "description": None,
        "notesModules": [],
        "requiredGames": [],
        "colorPresets": [{"name": name, "color": color} for color, name in palette.items()],
        "images": {},
        "tabs": tabs,
    }

    out = args.out or args.workbook.with_suffix(".json")
    out.write_text(json.dumps(checklist, indent=2), encoding="utf-8")
    total = sum(len(s["items"]) for t in tabs for s in t["sections"])
    print("\n{} items across {} tabs -> {}".format(total, len(tabs), out))
    return 0


if __name__ == "__main__":
    sys.exit(main())
